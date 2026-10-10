"""Prepara les imatges que fa servir genera.py. Només cal executar-lo una vegada.

- Retalla figures del cartell de la jornada i en fa el fons transparent.
- Treu les imatges del document «Elaborar material SIEI…» a conference/imatges/.
- Parteix en dues («abans» i «ara») les imatges del document que en tenen dues
  una sota l'altra, perquè cada una ocupi tota la pantalla.
- Fa imatges de pàgines de fitxes (les de Fast-build i les de 4t), sense marges.

    python3 conference/presentacio/prepara_recursos.py
"""
import io
from pathlib import Path

import docx
from docx.oxml.ns import qn
from PIL import Image

AQUI = Path(__file__).resolve().parent
CONF = AQUI.parent
REC = AQUI / "recursos"

# Figures del cartell: (esquerra, dalt, dreta, baix) en píxels del cartell de 853 × 1280
FIGURES = {
    "figura-1": (270, 415, 400, 585),   # noi en cadira de rodes amb un triangle
    "figura-2": (280, 670, 385, 870),   # dona amb bastó
    "figura-3": (690, 380, 830, 555),   # dues nenes, una amb un globus
    "figura-4": (560, 45, 780, 240),    # tres nens amb quadrats a la paret
    "figura-5": (620, 550, 800, 690),   # dos nens amb cubs
    "figura-6": (600, 690, 830, 840),   # dues persones a la taula, una en cadira
    "figura-7": (380, 680, 600, 880),   # grup que conversa
    "figura-8": (610, 860, 840, 1020),  # nens amb l'esquelet d'un cub
    "portada": (480, 20, 853, 1040),    # la columna dreta del cartell
}

# Imatges del document: nom de la imatge dins del .docx → nom del fitxer
DOCUMENT = {
    "image1.png": "les-tres-etapes.png",
    "image10.png": "caixa-eines-sense-format-i-amb-format.png",
    "image3.png": "calculadora-16-i-23-setembre.png",
    "image7.png": "a-la-vida-de-cada-dia.png",
    "image5.png": "porta-entrada-23-i-29-setembre.png",
    "image11.png": "taula-multiplicar-1r.png",
    "image4.png": "parabola-u5.png",
    "image2.png": "lletra-manuscrita-ud1.png",
    "image8.png": "regla-trencada-u3.png",
    "image6.png": "ud1-4t-pagina-a-pagina.png",
    "image9.png": "ud1-1r-fitxa-a-fitxa.png",
}


# Imatges del document amb l'«abans» a dalt i l'«ara» a sota
ABANS_ARA = ["calculadora-16-i-23-setembre.png", "caixa-eines-sense-format-i-amb-format.png"]

# Pàgines de fitxes: nom de la imatge → (PDF, pàgina començant per 1)
FITXES = {
    "fitxa-A-unitat1.png": ("Fast-build-A/Fast-build-A-Matematiques_Adaptades_4tESO.pdf", 2),
    "fitxa-B-equacions.png": ("Fast-build-B/3_Quadern_Algebra_i_funcions.pdf", 4),
    "fitxa-4t-equacions.png": ("../4eso/pdf/ud4-alumnat.pdf", 1),
}


def parteix_abans_ara(ruta):
    """Talla on comença la franja blava de l'«ARA» i desa les dues meitats."""
    im = Image.open(ruta).convert("RGB")
    tall = next(y for y in range(im.height // 4, im.height)
                if im.getpixel((4, y)) == (220, 235, 255))
    for nom, caixa in (("abans", (0, 0, im.width, tall)), ("ara", (0, tall, im.width, im.height))):
        desti = ruta.with_name(f"{ruta.stem}-{nom}.png")
        im.crop(caixa).save(desti)
        print("imatges/" + desti.name)


def sense_marges(im, marge=20):
    """Retalla el blanc que envolta el contingut i hi deixa un marge petit."""
    gris = im.convert("L").point(lambda v: 255 if v < 245 else 0)
    x0, y0, x1, y1 = gris.getbbox()
    return im.crop((max(0, x0 - marge), max(0, y0 - marge),
                    min(im.width, x1 + marge), min(im.height, y1 + marge)))


def fes_fitxes(sortida):
    import pypdfium2 as pdfium
    for nom, (pdf, pagina) in FITXES.items():
        doc = pdfium.PdfDocument(CONF / pdf)
        im = doc[pagina - 1].render(scale=150 / 72).to_pil().convert("RGB")
        sense_marges(im).save(sortida / nom)
        print("imatges/" + nom)


def transparent(im):
    """El fons gairebé blanc del cartell passa a ser transparent, amb la vora suau."""
    im = im.convert("RGBA")
    px = im.load()
    for y in range(im.height):
        for x in range(im.width):
            r, g, b, _ = px[x, y]
            d = 251 - min(r, g, b)
            px[x, y] = (r, g, b, max(0, min(255, d * 10)))
    return im


def main():
    cartell = Image.open(REC / "cartell.jpg").convert("RGB")
    for nom, caixa in FIGURES.items():
        im = cartell.crop(caixa)
        im = im.resize((im.width * 2, im.height * 2), Image.LANCZOS)
        transparent(im).save(REC / f"{nom}.png")
        print("recursos/" + nom + ".png")

    sortida = CONF / "imatges"
    sortida.mkdir(exist_ok=True)
    doc = docx.Document(CONF / "Elaborar material SIEI usant la IA i donant-li instruccions adequades.docx")
    for p in doc.paragraphs:
        for blip in p._p.iter(qn("a:blip")):
            part = doc.part.related_parts[blip.get(qn("r:embed"))]
            nom = DOCUMENT.get(Path(str(part.partname)).name)
            if nom:
                Image.open(io.BytesIO(part.blob)).save(sortida / nom)
                print("imatges/" + nom)
    for nom in ABANS_ARA:
        parteix_abans_ara(sortida / nom)
    fes_fitxes(sortida)


if __name__ == "__main__":
    main()
